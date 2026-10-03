from pathlib import Path

import pytest
from docx import Document

from axioms.llm import DisabledProvider, FakeProvider
from axioms.writing_agent import DocumentType, WritingRequest, build_writing_draft, export_docx


def request(document_type: DocumentType = DocumentType.PAPER_SECTION) -> WritingRequest:
    return WritingRequest(
        document_type=document_type,
        subject="Fuzzy similarity methods for drug-drug interaction prediction",
        audience="Q1-journal readers",
        purpose="Introduce the motivation and define the contribution boundaries.",
        key_points=(
            "Describe the practical motivation.",
            "State the method boundary without claiming unverified results.",
        ),
        verified_facts=(
            "The author has approved the problem statement.",
            "All citations will be checked before inclusion.",
        ),
        references=("Author-verified reference record pending insertion.",),
    )


def test_paper_section_draft_marks_evidence_for_verification() -> None:
    draft = build_writing_draft(request())
    assert draft.approval_required
    assert "Evidence placeholders" in [section.heading for section in draft.sections]
    assert "similarity score" in " ".join(draft.originality_checklist).casefold()


def test_email_draft_has_a_subject_line() -> None:
    draft = build_writing_draft(request(DocumentType.EMAIL))
    assert draft.subject_line == "Regarding: Fuzzy similarity methods for drug-drug interaction prediction"


def test_disabled_llm_retains_the_deterministic_review_first_draft() -> None:
    draft = build_writing_draft(request(), provider=DisabledProvider())

    assert draft.sections[0].content == "Introduce the motivation and define the contribution boundaries."
    assert draft.approval_required


def test_fake_llm_enriches_section_prose_without_changing_headings_or_review_gates() -> None:
    response = "\n---\n".join(f"Custom review-first paragraph {index}." for index in range(1, 6))
    captured_messages = []
    provider = FakeProvider(lambda messages: (captured_messages.extend(messages), response)[1])

    draft = build_writing_draft(request(), provider=provider)

    assert [section.heading for section in draft.sections] == [
        "Section purpose",
        "Author-verified foundation",
        "Proposed argument sequence",
        "Evidence placeholders",
        "Transition",
    ]
    assert draft.sections[0].content == "Custom review-first paragraph 1."
    assert draft.sections[-1].content == "Custom review-first paragraph 5."
    assert draft.approval_required
    assert "The author has approved the problem statement." in captured_messages[1].content
    assert "Verified facts (the only factual basis):" in captured_messages[1].content


def test_malformed_or_failed_llm_output_falls_back_to_the_template() -> None:
    malformed = build_writing_draft(request(), provider=FakeProvider(lambda _messages: "unstructured response"))

    def fail(_messages) -> str:
        raise RuntimeError("provider unavailable")

    failed = build_writing_draft(request(), provider=FakeProvider(fail))

    assert malformed.sections[0].content == "Introduce the motivation and define the contribution boundaries."
    assert failed.sections[0].content == "Introduce the motivation and define the contribution boundaries."


def test_recommendation_requires_two_verified_facts() -> None:
    invalid = WritingRequest(
        document_type=DocumentType.RECOMMENDATION_LETTER,
        subject="Candidate",
        audience="Selection committee",
        purpose="Provide a recommendation.",
        key_points=("Describe achievement.",),
        verified_facts=("One verified fact.",),
    )
    with pytest.raises(ValueError, match="at least two"):
        build_writing_draft(invalid)


def test_docx_export_contains_subject(tmp_path: Path) -> None:
    destination = export_docx(build_writing_draft(request()), tmp_path / "writing_draft.docx")
    document = Document(destination)
    assert destination.exists()
    assert "Fuzzy similarity methods" in document.paragraphs[0].text

