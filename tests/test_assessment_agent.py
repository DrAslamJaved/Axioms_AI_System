from pathlib import Path

import pytest
from docx import Document

from axioms.assessment_agent import (
    AssessmentRequest,
    AssessmentType,
    BloomLevel,
    Difficulty,
    LearningOutcome,
    build_assessment_blueprint,
    export_instructor_docx,
    export_student_docx,
)
from axioms.llm import DisabledProvider, FakeProvider


def request() -> AssessmentRequest:
    return AssessmentRequest(
        topic="Spectral Graph Theory",
        course_level="Graduate",
        assessment_type=AssessmentType.QUIZ,
        duration_minutes=30,
        total_marks=20,
        question_count=2,
        learning_outcomes=(
            LearningOutcome("LO1", "Explain the role of graph matrices.", BloomLevel.UNDERSTAND),
            LearningOutcome("LO2", "Apply a spectral method to a small graph.", BloomLevel.APPLY),
        ),
        approved_source_scope="Instructor-approved lecture notes on graph matrices and spectral methods.",
        difficulties=(Difficulty.EASY, Difficulty.MODERATE),
    )


def test_blueprint_allocates_all_marks_and_maps_outcomes() -> None:
    blueprint = build_assessment_blueprint(request())
    assert blueprint.allocated_marks == 20
    assert [question.outcome_id for question in blueprint.questions] == ["LO1", "LO2"]
    assert all(question.required_evidence for question in blueprint.questions)


def test_disabled_llm_retains_the_deterministic_instructor_blueprint() -> None:
    blueprint = build_assessment_blueprint(request(), provider=DisabledProvider())

    assert "Create one single-part, instructor-reviewed item" in blueprint.questions[0].prompt_framework
    assert blueprint.allocated_marks == 20


def test_fake_llm_enriches_internal_frameworks_without_changing_assessment_controls() -> None:
    response = "\n---\n".join(
        f"Internal instructor planning note {index} preserves the approved source boundary."
        for index in range(1, 3)
    )
    captured_messages = []
    provider = FakeProvider(lambda messages: (captured_messages.extend(messages), response)[1])

    blueprint = build_assessment_blueprint(request(), provider=provider)

    assert [question.marks for question in blueprint.questions] == [10, 10]
    assert [question.outcome_id for question in blueprint.questions] == ["LO1", "LO2"]
    assert blueprint.questions[0].prompt_framework.startswith("Internal instructor planning note 1")
    assert blueprint.instructor_review_required
    assert "Instructor-approved lecture notes" in captured_messages[1].content
    assert "Learning outcomes:" in captured_messages[1].content


def test_malformed_or_failed_llm_output_retains_the_instructor_blueprint() -> None:
    malformed = build_assessment_blueprint(request(), provider=FakeProvider(lambda _messages: "unstructured response"))

    def fail(_messages) -> str:
        raise RuntimeError("provider unavailable")

    failed = build_assessment_blueprint(request(), provider=FakeProvider(fail))

    assert "Create one single-part, instructor-reviewed item" in malformed.questions[0].prompt_framework
    assert "Create one single-part, instructor-reviewed item" in failed.questions[0].prompt_framework


def test_quiz_blueprint_rejects_more_than_two_items() -> None:
    invalid = AssessmentRequest(
        topic="Topic",
        course_level="Undergraduate",
        assessment_type=AssessmentType.QUIZ,
        duration_minutes=30,
        total_marks=30,
        question_count=3,
        learning_outcomes=(LearningOutcome("LO1", "Explain the concept.", BloomLevel.UNDERSTAND),),
        approved_source_scope="Approved lecture notes.",
    )
    with pytest.raises(ValueError, match="at most two"):
        build_assessment_blueprint(invalid)


def test_student_export_has_no_rubric_or_solution_guide(tmp_path: Path) -> None:
    destination = export_student_docx(build_assessment_blueprint(request()), tmp_path / "student.docx")
    text = "\n".join(paragraph.text for paragraph in Document(destination).paragraphs).casefold()
    assert "rubric" not in text
    assert "solution guide" not in text


def test_instructor_export_includes_rubric(tmp_path: Path) -> None:
    destination = export_instructor_docx(build_assessment_blueprint(request()), tmp_path / "instructor.docx")
    text = "\n".join(paragraph.text for paragraph in Document(destination).paragraphs).casefold()
    assert "rubric" in text

