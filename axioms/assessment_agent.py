"""Instructor-facing, assessment-blueprint agent with explicit academic-integrity controls."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document


class AssessmentType(StrEnum):
    QUIZ = "quiz"
    ASSIGNMENT = "assignment"
    CLASS_ACTIVITY = "class_activity"
    MIDTERM = "midterm"
    FINAL = "final"


class BloomLevel(StrEnum):
    REMEMBER = "remember"
    UNDERSTAND = "understand"
    APPLY = "apply"
    ANALYZE = "analyze"
    EVALUATE = "evaluate"
    CREATE = "create"


class Difficulty(StrEnum):
    EASY = "easy"
    MODERATE = "moderate"
    HARD = "hard"


VERBS = {
    BloomLevel.REMEMBER: "identify and state",
    BloomLevel.UNDERSTAND: "explain and interpret",
    BloomLevel.APPLY: "apply and justify",
    BloomLevel.ANALYZE: "analyze and distinguish",
    BloomLevel.EVALUATE: "evaluate and defend",
    BloomLevel.CREATE: "design and justify",
}


@dataclass(frozen=True, slots=True)
class LearningOutcome:
    outcome_id: str
    text: str
    bloom_level: BloomLevel

    def validate(self) -> None:
        if not self.outcome_id.strip() or not self.text.strip():
            raise ValueError("Each learning outcome needs an ID and observable text.")


@dataclass(frozen=True, slots=True)
class AssessmentRequest:
    topic: str
    course_level: str
    assessment_type: AssessmentType
    duration_minutes: int
    total_marks: int
    question_count: int
    learning_outcomes: tuple[LearningOutcome, ...]
    approved_source_scope: str
    difficulties: tuple[Difficulty, ...] = (Difficulty.MODERATE,)
    require_handwritten_work: bool = True
    require_reflection: bool = True
    include_personalised_context: bool = True

    def validate(self) -> None:
        if not self.topic.strip() or not self.course_level.strip() or not self.approved_source_scope.strip():
            raise ValueError("Topic, course level, and approved source scope are required.")
        if not 10 <= self.duration_minutes <= 240:
            raise ValueError("Assessment duration must be between 10 and 240 minutes.")
        if not 1 <= self.question_count <= 20:
            raise ValueError("Question count must be between one and twenty.")
        if self.total_marks < self.question_count:
            raise ValueError("Total marks must be at least the number of questions.")
        if not 1 <= len(self.learning_outcomes) <= 12:
            raise ValueError("Provide between one and twelve learning outcomes.")
        if not self.difficulties:
            raise ValueError("Provide at least one difficulty level.")
        outcome_ids = [outcome.outcome_id for outcome in self.learning_outcomes]
        if len(outcome_ids) != len(set(outcome_ids)):
            raise ValueError("Learning-outcome IDs must be unique.")
        for outcome in self.learning_outcomes:
            outcome.validate()
        if self.assessment_type is AssessmentType.QUIZ and self.question_count > 2:
            raise ValueError("This MVP uses at most two single-part items for a quiz blueprint.")


@dataclass(frozen=True, slots=True)
class RubricCriterion:
    criterion: str
    weight_percent: int
    descriptor: str


@dataclass(frozen=True, slots=True)
class QuestionBlueprint:
    number: int
    outcome_id: str
    bloom_level: BloomLevel
    difficulty: Difficulty
    marks: int
    prompt_framework: str
    required_evidence: tuple[str, ...]
    rubric: tuple[RubricCriterion, ...]


@dataclass(frozen=True, slots=True)
class AssessmentBlueprint:
    request: AssessmentRequest
    questions: tuple[QuestionBlueprint, ...]
    ai_resilience_review: tuple[str, ...]
    quality_checks: tuple[str, ...]
    instructor_review_required: bool = True

    @property
    def allocated_marks(self) -> int:
        return sum(question.marks for question in self.questions)

    def to_dict(self) -> dict:
        return asdict(self)


def _allocate_marks(total: int, count: int) -> tuple[int, ...]:
    base, extra = divmod(total, count)
    return tuple(base + (1 if index < extra else 0) for index in range(count))


def _required_evidence(request: AssessmentRequest) -> tuple[str, ...]:
    evidence = ["Show the method, assumptions, intermediate steps, and a final interpretation."]
    if request.require_handwritten_work:
        evidence.append("Provide handwritten calculations, an annotated diagram, or equivalent working evidence.")
    if request.include_personalised_context:
        evidence.append("Respond to the instructor-assigned context, data values, or scenario variant.")
    if request.require_reflection:
        evidence.append("Add a brief reflection on a decision, limitation, or likely error source.")
    return tuple(evidence)


def _rubric() -> tuple[RubricCriterion, ...]:
    return (
        RubricCriterion("Method and reasoning", 40, "Valid method, assumptions, and justified intermediate steps."),
        RubricCriterion("Correctness", 30, "Accurate calculations, notation, and final result where applicable."),
        RubricCriterion("Interpretation", 20, "Meaningful explanation of the result in the assigned context."),
        RubricCriterion("Academic integrity and presentation", 10, "Clear work, required evidence, and authentic individual response."),
    )


def build_assessment_blueprint(request: AssessmentRequest) -> AssessmentBlueprint:
    """Build a source-bounded instructor blueprint, not an autonomous student assessment."""
    request.validate()
    marks = _allocate_marks(request.total_marks, request.question_count)
    questions: list[QuestionBlueprint] = []
    for index, mark in enumerate(marks):
        outcome = request.learning_outcomes[index % len(request.learning_outcomes)]
        difficulty = request.difficulties[index % len(request.difficulties)]
        verb = VERBS[outcome.bloom_level]
        prompt = (
            f"Create one single-part, instructor-reviewed item on '{request.topic}' that asks students to "
            f"{verb} the approved course material for {outcome.outcome_id}. Use only this approved scope: "
            f"{request.approved_source_scope}"
        )
        questions.append(
            QuestionBlueprint(
                number=index + 1,
                outcome_id=outcome.outcome_id,
                bloom_level=outcome.bloom_level,
                difficulty=difficulty,
                marks=mark,
                prompt_framework=prompt,
                required_evidence=_required_evidence(request),
                rubric=_rubric(),
            )
        )
    ai_review = [
        "Require reasoning and intermediate work; do not assess a final answer alone.",
        "Use a course-specific, instructor-approved scenario rather than a directly searchable wording.",
        "Require a short oral, written, or visual explanation that connects the method to the assigned context.",
        "Treat this as a qualitative design review, not a claim that an assessment is AI-proof.",
    ]
    if request.include_personalised_context:
        ai_review.append("Create parallel versions by changing approved values or contexts while preserving the blueprint.")
    quality_checks = (
        "Verify all mathematical statements, data values, and notation against approved course sources.",
        "Confirm each item maps to exactly the listed learning outcome and stated Bloom level.",
        "Check that marks, timing, and difficulty are appropriate for the actual class context.",
        "Keep solution guides and rubrics in the instructor package; never include them in student-facing material.",
        "Do not store student names, grades, submissions, or other personal data in this agent.",
        "Obtain instructor approval before printing, LMS upload, release, marking, or feedback use.",
    )
    return AssessmentBlueprint(
        request=request,
        questions=tuple(questions),
        ai_resilience_review=tuple(ai_review),
        quality_checks=quality_checks,
    )


def export_student_docx(blueprint: AssessmentBlueprint, destination: Path) -> Path:
    """Export a student-facing blueprint without rubrics, solutions, or internal controls."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"Assessment: {blueprint.request.topic}", level=0)
    document.add_paragraph(
        f"{blueprint.request.assessment_type.value.title()} | {blueprint.request.duration_minutes} minutes | "
        f"{blueprint.request.total_marks} marks"
    )
    document.add_paragraph("Complete every item using the format and evidence specified by your instructor.")
    for question in blueprint.questions:
        document.add_heading(f"Question {question.number} ({question.marks} marks)", level=1)
        document.add_paragraph("Instructor-approved question text will be inserted before release.")
        document.add_paragraph("Required response evidence:")
        for item in question.required_evidence:
            document.add_paragraph(item, style="List Bullet")
    document.save(destination)
    return destination


def export_instructor_docx(blueprint: AssessmentBlueprint, destination: Path) -> Path:
    """Export internal blueprint, rubrics, and quality checks for instructor review."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"Instructor Assessment Blueprint: {blueprint.request.topic}", level=0)
    document.add_paragraph(f"Approved source scope: {blueprint.request.approved_source_scope}")
    document.add_heading("Learning outcomes", level=1)
    for outcome in blueprint.request.learning_outcomes:
        document.add_paragraph(f"{outcome.outcome_id} — {outcome.bloom_level.value}: {outcome.text}", style="List Bullet")
    for question in blueprint.questions:
        document.add_heading(f"Question {question.number}: {question.marks} marks", level=1)
        document.add_paragraph(question.prompt_framework)
        document.add_paragraph(
            f"Blueprint: {question.outcome_id}; {question.bloom_level.value}; {question.difficulty.value}"
        )
        document.add_paragraph("Rubric:")
        for criterion in question.rubric:
            document.add_paragraph(
                f"{criterion.criterion} ({criterion.weight_percent}%): {criterion.descriptor}", style="List Bullet"
            )
    document.add_heading("AI-resilience review", level=1)
    for item in blueprint.ai_resilience_review:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Quality checks before release", level=1)
    for item in blueprint.quality_checks:
        document.add_paragraph(item, style="List Bullet")
    document.save(destination)
    return destination

