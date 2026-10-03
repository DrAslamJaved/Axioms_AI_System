"""Deterministic, review-first lecture-design agent for the Axioms MVP."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document

from axioms.llm import LLMDisabledError, LLMMessage, LLMProvider, get_provider


@dataclass(frozen=True, slots=True)
class LectureRequest:
    topic: str
    course_level: str
    duration_minutes: int
    audience: str
    prior_knowledge: str = "Not yet specified"
    learning_outcomes: tuple[str, ...] = ()
    application_context: str | None = None
    include_computational_activity: bool = False

    def validate(self) -> None:
        if not self.topic.strip():
            raise ValueError("A lecture topic is required.")
        if not 30 <= self.duration_minutes <= 240:
            raise ValueError("Lecture duration must be between 30 and 240 minutes.")
        if not 2 <= len(self.learning_outcomes) <= 6:
            raise ValueError("Provide between two and six learning outcomes.")
        if any(not outcome.strip() for outcome in self.learning_outcomes):
            raise ValueError("Learning outcomes cannot be blank.")


@dataclass(frozen=True, slots=True)
class TimedSection:
    title: str
    minutes: int
    purpose: str
    instructor_prompt: str


@dataclass(frozen=True, slots=True)
class LecturePlan:
    request: LectureRequest
    sections: tuple[TimedSection, ...]
    board_work: tuple[str, ...]
    practice_activity: str
    review_checklist: tuple[str, ...]

    @property
    def total_minutes(self) -> int:
        return sum(section.minutes for section in self.sections)

    def to_dict(self) -> dict:
        return asdict(self)

    def to_markdown(self) -> str:
        request = self.request
        rows = "\n".join(
            f"| {section.minutes} | {section.title} | {section.purpose} |"
            for section in self.sections
        )
        outcomes = "\n".join(f"- {outcome}" for outcome in request.learning_outcomes)
        board_work = "\n".join(f"- {item}" for item in self.board_work)
        review = "\n".join(f"- [ ] {item}" for item in self.review_checklist)
        return f"""# Lecture plan: {request.topic}

## Teaching context
- Course level: {request.course_level}
- Audience: {request.audience}
- Duration: {request.duration_minutes} minutes
- Prior knowledge: {request.prior_knowledge}
- Application context: {request.application_context or 'To be selected by the instructor'}

## Learning outcomes
{outcomes}

## Timed teaching plan
| Minutes | Segment | Purpose |
| ---: | --- | --- |
{rows}

## Board-work / slide cues
{board_work}

## Guided practice
{self.practice_activity}

## Instructor review before use
{review}
"""


def _allocate_minutes(total: int) -> tuple[int, ...]:
    """Allocate a fixed duration over the teaching sequence without rounding drift."""
    weights = (0.10, 0.18, 0.27, 0.25, 0.15, 0.05)
    raw = [total * weight for weight in weights]
    allocation = [max(1, int(value)) for value in raw]
    remainder = total - sum(allocation)
    fractions = sorted(
        range(len(raw)), key=lambda index: (raw[index] - int(raw[index]), -index), reverse=True
    )
    for index in fractions[:remainder]:
        allocation[index] += 1
    return tuple(allocation)


def build_lecture_plan(request: LectureRequest, *, provider: LLMProvider | None = None) -> LecturePlan:
    """Build a review-first plan, enriching its fixed teaching structure only when configured."""
    request.validate()
    minutes = _allocate_minutes(request.duration_minutes)
    application = request.application_context or "a discipline-relevant real-world example"
    computational_activity = (
        "Then implement or explore one small example in Python/NumPy; ask students to explain "
        "how the output connects to the formal method."
        if request.include_computational_activity
        else "Then ask students to compare two solution approaches and justify their choice."
    )
    sections = (
        TimedSection("Opening hook", minutes[0], "Activate curiosity through a familiar problem.", f"Where might {request.topic} arise outside the textbook?"),
        TimedSection("Intuition and prior knowledge", minutes[1], "Connect the idea to what students already know.", "Ask for a prediction before introducing notation."),
        TimedSection("Formal development", minutes[2], "State definitions, assumptions, notation, and theorem conditions precisely.", "Pause after each definition: what does each symbol mean?"),
        TimedSection("Worked example", minutes[3], "Model a complete solution with reasoning visible.", "Ask students to identify the next valid step before you reveal it."),
        TimedSection("Guided application", minutes[4], f"Apply the method to {application}.", "Circulate, prompt, and collect one representative misconception."),
        TimedSection("Retrieval and exit check", minutes[5], "Consolidate learning and expose remaining misconceptions.", "Use a one-minute written explanation or a short formative question."),
    )
    sections = _synthesise_sections(request, sections, provider or get_provider())
    board_work = (
        f"Start with a labelled motivating situation for {request.topic}; do not start with symbols.",
        "Build notation progressively and box assumptions before using them.",
        "Leave a visible 'common misconception' note beside the worked example.",
        "Finish with a compact concept map linking intuition, definition, method, and application.",
    )
    review_checklist = (
        "Verify every theorem statement, calculation, and example against the approved course source.",
        "Align the plan with the course outline and the actual class duration.",
        "Confirm learning outcomes use observable action verbs and an appropriate Bloom level.",
        "Check accessibility, mathematical notation, and cultural relevance of the opening example.",
        "Do not include student names, grades, or other personal data.",
    )
    return LecturePlan(
        request=request,
        sections=sections,
        board_work=board_work,
        practice_activity=(
            f"In pairs, solve a short {request.topic} problem, annotate the assumption used at each step, "
            f"and connect the result to {application}. {computational_activity}"
        ),
        review_checklist=review_checklist,
    )


_LECTURE_SYSTEM_PROMPT = """You are the Axioms lecture-design assistant.
Improve only the pedagogical purpose and instructor prompt for each supplied timed
section. Preserve the requested duration, section order, learning outcomes, and
the requirement for human instructor review. Do not assert unverified theorems,
facts, calculations, examples, or citations. Do not include student data.

Return exactly one block per supplied section, in its original order, separated
by a line containing only `---`. Each block must contain exactly these two lines:
Purpose: <one concise sentence>
Instructor prompt: <one concise sentence>
"""


def _synthesise_sections(
    request: LectureRequest,
    sections: tuple[TimedSection, ...],
    provider: LLMProvider,
) -> tuple[TimedSection, ...]:
    """Use bounded pedagogical wording when available; retain the deterministic plan otherwise."""
    section_specification = "\n".join(
        f"{index}. {section.title} ({section.minutes} minutes)"
        for index, section in enumerate(sections, start=1)
    )
    outcomes = "\n".join(f"- {outcome}" for outcome in request.learning_outcomes)
    user_prompt = (
        f"Topic: {request.topic}\n"
        f"Course level: {request.course_level}\n"
        f"Audience: {request.audience}\n"
        f"Prior knowledge: {request.prior_knowledge}\n"
        f"Application context: {request.application_context or 'Not specified'}\n"
        f"Learning outcomes:\n{outcomes}\n\n"
        f"Timed sections:\n{section_specification}\n"
    )
    try:
        result = provider.complete(
            [LLMMessage("system", _LECTURE_SYSTEM_PROMPT), LLMMessage("user", user_prompt)],
            max_tokens=900,
            temperature=0.2,
        )
    except LLMDisabledError:
        return sections
    except Exception:  # noqa: BLE001 - provider failures must retain the local review-first fallback
        return sections
    generated = _parse_section_synthesis(result.text, len(sections))
    if generated is None:
        return sections
    return tuple(
        TimedSection(
            title=section.title,
            minutes=section.minutes,
            purpose=purpose,
            instructor_prompt=instructor_prompt,
        )
        for section, (purpose, instructor_prompt) in zip(sections, generated, strict=True)
    )


def _parse_section_synthesis(text: str, expected_sections: int) -> tuple[tuple[str, str], ...] | None:
    """Accept only the narrow response format; malformed model output never alters the plan."""
    blocks = [block.strip() for block in text.split("---") if block.strip()]
    if len(blocks) != expected_sections:
        return None
    parsed: list[tuple[str, str]] = []
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if len(lines) != 2 or not lines[0].startswith("Purpose: ") or not lines[1].startswith("Instructor prompt: "):
            return None
        purpose = lines[0].removeprefix("Purpose: ").strip()
        instructor_prompt = lines[1].removeprefix("Instructor prompt: ").strip()
        if not purpose or not instructor_prompt or len(purpose) > 600 or len(instructor_prompt) > 600:
            return None
        parsed.append((purpose, instructor_prompt))
    return tuple(parsed)


def export_docx(plan: LecturePlan, destination: Path) -> Path:
    """Export a reviewed lesson plan. The caller owns the output directory and file lifecycle."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"Lecture Plan: {plan.request.topic}", level=0)
    document.add_paragraph(
        f"{plan.request.course_level} | {plan.request.audience} | {plan.total_minutes} minutes"
    )
    document.add_heading("Learning outcomes", level=1)
    for outcome in plan.request.learning_outcomes:
        document.add_paragraph(outcome, style="List Bullet")
    document.add_heading("Timed teaching plan", level=1)
    table = document.add_table(rows=1, cols=3)
    for cell, heading in zip(table.rows[0].cells, ("Minutes", "Segment", "Purpose"), strict=True):
        cell.text = heading
    for section in plan.sections:
        cells = table.add_row().cells
        cells[0].text = str(section.minutes)
        cells[1].text = section.title
        cells[2].text = section.purpose
    document.add_heading("Board-work / slide cues", level=1)
    for cue in plan.board_work:
        document.add_paragraph(cue, style="List Bullet")
    document.add_heading("Guided practice", level=1)
    document.add_paragraph(plan.practice_activity)
    document.add_heading("Instructor review before use", level=1)
    for item in plan.review_checklist:
        document.add_paragraph(item, style="List Bullet")
    document.save(destination)
    return destination

