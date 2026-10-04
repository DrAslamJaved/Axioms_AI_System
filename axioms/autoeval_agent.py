"""Deterministic, review-first quality checks for Axioms agent deliverables."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from enum import StrEnum
from hashlib import sha256
from pathlib import Path

from docx import Document

from axioms.llm import LLMDisabledError, LLMMessage, LLMProvider


class EvaluatedAgent(StrEnum):
    LECTURE = "lecture_design"
    WRITING = "writing_communication"
    RESEARCH = "research"
    ASSESSMENT = "assessment_design"
    CONTENT = "content_creation"
    SOCIAL_MEDIA = "social_media"
    PORTFOLIO = "stem_ai_portfolio"


class CheckStatus(StrEnum):
    PASS = "pass"
    REVIEW = "review"
    BLOCK = "block"


@dataclass(frozen=True, slots=True)
class AutoEvalRequest:
    evaluated_agent: EvaluatedAgent
    deliverable_title: str
    artifact_text: str
    required_elements: tuple[str, ...]
    evidence_markers: tuple[str, ...] = ()
    public_facing: bool = False
    declared_sensitive_data: bool = False

    def validate(self) -> None:
        if not self.deliverable_title.strip() or not self.artifact_text.strip():
            raise ValueError("Deliverable title and artifact text are required.")
        if not 1 <= len(self.required_elements) <= 20:
            raise ValueError("Provide between one and twenty required elements.")
        if any(not item.strip() for item in self.required_elements):
            raise ValueError("Required elements cannot be blank.")
        if len(set(self.required_elements)) != len(self.required_elements):
            raise ValueError("Required elements must be unique.")
        if any(not item.strip() for item in self.evidence_markers):
            raise ValueError("Evidence markers cannot be blank.")


@dataclass(frozen=True, slots=True)
class EvaluationCheck:
    name: str
    status: CheckStatus
    detail: str


@dataclass(frozen=True, slots=True)
class AutoEvalReport:
    request: AutoEvalRequest
    artifact_sha256: str
    checks: tuple[EvaluationCheck, ...]
    required_elements_found: tuple[str, ...]
    missing_required_elements: tuple[str, ...]
    quality_signal_percent: int
    review_boundary: str
    qualitative_summary: str | None = None
    human_review_required: bool = True
    automatic_reconfiguration_blocked: bool = True
    external_action_blocked: bool = True

    def to_dict(self) -> dict:
        return asdict(self)

    def to_markdown(self) -> str:
        rows = "\n".join(
            f"| {item.name} | {item.status.value} | {item.detail} |" for item in self.checks
        )
        found = "\n".join(f"- {item}" for item in self.required_elements_found) or "- None."
        missing = "\n".join(f"- {item}" for item in self.missing_required_elements) or "- None."
        return f"""# AutoEval report: {self.request.deliverable_title}

## Audit identity
- Evaluated agent: {self.request.evaluated_agent.value}
- Artifact SHA-256: `{self.artifact_sha256}`
- Quality signal: {self.quality_signal_percent}%

## Deterministic checks
| Check | Status | Detail |
| --- | --- | --- |
{rows}

## Required elements found
{found}

## Required elements needing revision
{missing}

## Review boundary
{self.review_boundary}

## Optional qualitative summary
{self.qualitative_summary or "No provider-generated qualitative summary is available; use the deterministic checks above."}
"""


def _contains(text: str, item: str) -> bool:
    return item.casefold() in text.casefold()


_AUTOEVAL_SYSTEM_PROMPT = """You are the Axioms internal AutoEval summariser.
Write only one concise qualitative summary of the supplied deterministic check
results. The deterministic SHA-256, statuses, required-marker results, score,
and review boundary are authoritative; you may not modify, reinterpret, replace,
or add to them. Do not claim factual, citation, originality, accessibility,
pedagogical, legal, privacy, or release verification. Do not approve, reject,
reconfigure, publish, schedule, message, or authorize anything. This is a
non-authoritative review note for a human reviewer.

Return one paragraph only, with no heading, list, citation, link, or question.
"""


def _synthesise_qualitative_summary(
    request: AutoEvalRequest,
    checks: tuple[EvaluationCheck, ...],
    found: tuple[str, ...],
    missing: tuple[str, ...],
    quality_signal: int,
    review_boundary: str,
    provider: LLMProvider | None,
) -> str | None:
    """Keep deterministic evaluation authoritative; use a provider only for a bounded human-review note."""
    if provider is None:
        return None
    check_results = "\n".join(
        f"- {check.name}: {check.status.value}; {check.detail}" for check in checks
    )
    user_prompt = (
        f"Evaluated agent: {request.evaluated_agent.value}\n"
        f"Deliverable title: {request.deliverable_title}\n"
        f"Artifact SHA-256: {sha256(request.artifact_text.encode('utf-8')).hexdigest()}\n"
        f"Quality signal: {quality_signal}%\n"
        f"Required elements found: {', '.join(found) or 'none'}\n"
        f"Required elements missing: {', '.join(missing) or 'none'}\n\n"
        f"Deterministic check results:\n{check_results}\n\n"
        f"Review boundary:\n{review_boundary}\n"
    )
    try:
        result = provider.complete(
            [LLMMessage("system", _AUTOEVAL_SYSTEM_PROMPT), LLMMessage("user", user_prompt)],
            max_tokens=500,
            temperature=0.0,
        )
    except LLMDisabledError:
        return None
    except Exception:  # noqa: BLE001 - provider failures must never alter deterministic QA
        return None
    summary = result.text.strip()
    if (
        not summary
        or "\n" in summary
        or "?" in summary
        or "http://" in summary.casefold()
        or "https://" in summary.casefold()
        or summary.startswith(("-", "*", "#"))
        or len(summary) > 1_500
    ):
        return None
    return summary


def evaluate_deliverable(
    request: AutoEvalRequest, *, provider: LLMProvider | None = None
) -> AutoEvalReport:
    """Run lexical contract checks only; never claim semantic, factual, or pedagogical verification."""
    request.validate()
    found = tuple(item for item in request.required_elements if _contains(request.artifact_text, item))
    missing = tuple(item for item in request.required_elements if item not in found)
    checks: list[EvaluationCheck] = []
    if missing:
        checks.append(
            EvaluationCheck(
                "Required-element coverage",
                CheckStatus.REVIEW,
                "Missing required text markers: " + ", ".join(missing),
            )
        )
    else:
        checks.append(
            EvaluationCheck(
                "Required-element coverage",
                CheckStatus.PASS,
                "All requested text markers are present.",
            )
        )
    missing_evidence = tuple(
        marker for marker in request.evidence_markers if not _contains(request.artifact_text, marker)
    )
    if not request.evidence_markers:
        checks.append(
            EvaluationCheck(
                "Evidence-marker coverage",
                CheckStatus.REVIEW,
                "No evidence markers were supplied; factual accuracy was not evaluated.",
            )
        )
    elif missing_evidence:
        checks.append(
            EvaluationCheck(
                "Evidence-marker coverage",
                CheckStatus.REVIEW,
                "Missing evidence markers: " + ", ".join(missing_evidence),
            )
        )
    else:
        checks.append(
            EvaluationCheck(
                "Evidence-marker coverage",
                CheckStatus.PASS,
                "All supplied evidence markers are present; their truth is not independently verified.",
            )
        )
    if request.declared_sensitive_data:
        checks.append(
            EvaluationCheck(
                "Sensitive-data declaration",
                CheckStatus.BLOCK,
                "Declared sensitive data blocks release; remove or properly govern it before further review.",
            )
        )
    else:
        checks.append(
            EvaluationCheck(
                "Sensitive-data declaration",
                CheckStatus.PASS,
                "No sensitive data was declared; this is not an independent privacy scan.",
            )
        )
    if request.public_facing:
        checks.append(
            EvaluationCheck(
                "Public-delivery gate",
                CheckStatus.REVIEW,
                "Public-facing material requires explicit human approval; AutoEval cannot grant approval.",
            )
        )
    score_checks = [item for item in checks if item.status is CheckStatus.PASS]
    quality_signal = round(100 * len(score_checks) / len(checks))
    boundary = (
        "This report checks declared text markers and declarations only. It does not establish "
        "factual accuracy, citation validity, originality, accessibility, pedagogy, legal compliance, "
        "or suitability for release. Scores are review signals and never trigger automatic approval, "
        "agent reconfiguration, publication, scheduling, or external action."
    )
    deterministic_checks = tuple(checks)
    qualitative_summary = _synthesise_qualitative_summary(
        request,
        deterministic_checks,
        found,
        missing,
        quality_signal,
        boundary,
        provider,
    )
    return AutoEvalReport(
        request=request,
        artifact_sha256=sha256(request.artifact_text.encode("utf-8")).hexdigest(),
        checks=deterministic_checks,
        required_elements_found=found,
        missing_required_elements=missing,
        quality_signal_percent=quality_signal,
        review_boundary=boundary,
        qualitative_summary=qualitative_summary,
    )


def export_docx(report: AutoEvalReport, destination: Path) -> Path:
    """Export a human-review report; it makes no change to the evaluated agent or artifact."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"AutoEval Report: {report.request.deliverable_title}", level=0)
    document.add_paragraph(f"Evaluated agent: {report.request.evaluated_agent.value}")
    document.add_paragraph(f"Artifact SHA-256: {report.artifact_sha256}")
    document.add_paragraph(f"Quality signal: {report.quality_signal_percent}% (review signal only)")
    document.add_heading("Deterministic checks", level=1)
    table = document.add_table(rows=1, cols=3)
    for cell, heading in zip(table.rows[0].cells, ("Check", "Status", "Detail"), strict=True):
        cell.text = heading
    for item in report.checks:
        cells = table.add_row().cells
        cells[0].text = item.name
        cells[1].text = item.status.value
        cells[2].text = item.detail
    document.add_heading("Review boundary", level=1)
    document.add_paragraph(report.review_boundary)
    document.add_heading("Optional qualitative summary", level=1)
    document.add_paragraph(
        report.qualitative_summary
        or "No provider-generated qualitative summary is available; use the deterministic checks above."
    )
    document.save(destination)
    return destination
