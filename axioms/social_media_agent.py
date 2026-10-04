"""Human-governed social-media drafting for academic and educational communication."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document

from axioms.llm import LLMDisabledError, LLMMessage, LLMProvider, get_provider


class SocialPlatform(StrEnum):
    LINKEDIN = "linkedin"
    INSTAGRAM = "instagram"
    X = "x"
    TIKTOK = "tiktok"
    WHATSAPP = "whatsapp"


class SocialObjective(StrEnum):
    EDUCATE = "educate"
    ANNOUNCE = "announce"
    INVITE_DISCUSSION = "invite_discussion"


@dataclass(frozen=True, slots=True)
class RecentSocialPost:
    platform: SocialPlatform
    topic: str
    hours_since_publication: int

    def validate(self) -> None:
        if not self.topic.strip():
            raise ValueError("Recent-post topics cannot be blank.")
        if self.hours_since_publication < 0:
            raise ValueError("Hours since publication cannot be negative.")


@dataclass(frozen=True, slots=True)
class SocialMediaRequest:
    topic: str
    audience: str
    platforms: tuple[SocialPlatform, ...]
    objective: SocialObjective
    approved_source_scope: str
    verified_facts: tuple[str, ...]
    brand_voice: str = "clear, respectful, and evidence-aware"
    call_to_action: str | None = None
    calendar_weeks: int = 4
    recent_posts: tuple[RecentSocialPost, ...] = ()

    def validate(self) -> None:
        if not self.topic.strip() or not self.audience.strip() or not self.approved_source_scope.strip():
            raise ValueError("Topic, audience, and approved source scope are required.")
        if not 1 <= len(self.platforms) <= 5:
            raise ValueError("Choose between one and five social platforms.")
        if len(set(self.platforms)) != len(self.platforms):
            raise ValueError("Each social platform may be selected only once.")
        if not 1 <= len(self.verified_facts) <= 8 or any(not fact.strip() for fact in self.verified_facts):
            raise ValueError("Provide between one and eight non-blank verified facts.")
        if not self.brand_voice.strip():
            raise ValueError("Brand voice cannot be blank.")
        if not 1 <= self.calendar_weeks <= 4:
            raise ValueError("Calendar length must be between one and four weeks.")
        for post in self.recent_posts:
            post.validate()
            if (
                post.platform in self.platforms
                and _normalise_topic(post.topic) == _normalise_topic(self.topic)
                and post.hours_since_publication < 48
            ):
                raise ValueError(
                    f"The 48-hour cooldown applies to {post.platform.value} for this topic."
                )


@dataclass(frozen=True, slots=True)
class PlatformDraft:
    platform: SocialPlatform
    format: str
    headline: str
    draft_copy: str
    asset_brief: str
    accessibility_note: str


@dataclass(frozen=True, slots=True)
class CalendarEntry:
    week: int
    platform: SocialPlatform
    purpose: str
    status: str = "draft only — not scheduled"


@dataclass(frozen=True, slots=True)
class SocialMediaPackage:
    request: SocialMediaRequest
    platform_drafts: tuple[PlatformDraft, ...]
    proposed_calendar: tuple[CalendarEntry, ...]
    publication_checks: tuple[str, ...]
    approval_required: bool = True
    external_action_blocked: bool = True
    topic_cooldown_hours: int = 48

    def to_dict(self) -> dict:
        return asdict(self)

    def to_markdown(self) -> str:
        platform_sections = "\n\n".join(
            f"## {item.platform.value.title()} — {item.format}\n"
            f"### {item.headline}\n{item.draft_copy}\n\n"
            f"**Asset brief:** {item.asset_brief}\n\n"
            f"**Accessibility:** {item.accessibility_note}"
            for item in self.platform_drafts
        )
        calendar = "\n".join(
            f"| {item.week} | {item.platform.value} | {item.purpose} | {item.status} |"
            for item in self.proposed_calendar
        )
        checks = "\n".join(f"- [ ] {item}" for item in self.publication_checks)
        return f"""# Social media package: {self.request.topic}

## Context
- Audience: {self.request.audience}
- Objective: {self.request.objective.value.replace('_', ' ')}
- Approved source scope: {self.request.approved_source_scope}
- Voice: {self.request.brand_voice}

{platform_sections}

## Proposed calendar
| Week | Platform | Purpose | Status |
| ---: | --- | --- | --- |
{calendar}

## Publication checks
{checks}
"""


def _normalise_topic(topic: str) -> str:
    return " ".join(topic.casefold().split())


def _call_to_action(request: SocialMediaRequest) -> str:
    return request.call_to_action or "Invite thoughtful questions or a review of the approved source material."


def _draft_for_platform(request: SocialMediaRequest, platform: SocialPlatform) -> PlatformDraft:
    fact = request.verified_facts[0]
    cta = _call_to_action(request)
    if platform is SocialPlatform.LINKEDIN:
        return PlatformDraft(
            platform=platform,
            format="professional post",
            headline=f"A careful introduction to {request.topic}",
            draft_copy=(
                f"For {request.audience}: {fact}\n\n"
                f"This draft explains one focused aspect of {request.topic} for an {request.objective.value.replace('_', ' ')} purpose. "
                f"{cta}\n\nSource scope for author review: {request.approved_source_scope}"
            ),
            asset_brief="Use an original diagram or a licensed visual with a clear source attribution.",
            accessibility_note="Add meaningful alt text and retain readable line breaks before publication.",
        )
    if platform is SocialPlatform.INSTAGRAM:
        return PlatformDraft(
            platform=platform,
            format="five-slide carousel",
            headline=f"{request.topic}: a reviewable visual explainer",
            draft_copy=(
                "Slide 1: Accurate title and learning promise.\n"
                f"Slide 2: Verified context — {fact}\n"
                "Slide 3: One carefully checked definition or idea.\n"
                "Slide 4: A source-reviewed example or limitation.\n"
                f"Slide 5: {cta}\n\n"
                "Caption must cite or link the author-approved source scope before publication."
            ),
            asset_brief="Use large, high-contrast text and original or appropriately licensed diagrams.",
            accessibility_note="Provide alt text for every slide and do not place essential text only inside images.",
        )
    if platform is SocialPlatform.X:
        return PlatformDraft(
            platform=platform,
            format="five-post thread",
            headline=f"A concise thread on {request.topic}",
            draft_copy=(
                f"1/5 What is the precise question behind {request.topic}?\n"
                f"2/5 Verified context: {fact}\n"
                "3/5 State one checked definition, assumption, or limitation.\n"
                "4/5 Give a short source-reviewed example; avoid overclaiming.\n"
                f"5/5 {cta}"
            ),
            asset_brief="Attach only an original, licensed, or permission-cleared visual if one improves clarity.",
            accessibility_note="Describe any attached visual in the post text and avoid image-only explanations.",
        )
    if platform is SocialPlatform.TIKTOK:
        return PlatformDraft(
            platform=platform,
            format="short-video script outline",
            headline=f"{request.topic} in one accurate minute",
            draft_copy=(
                "0–5s: Ask a plain-language, non-sensational question.\n"
                f"5–20s: State verified context — {fact}\n"
                "20–40s: Explain one checked idea with an original visual.\n"
                "40–55s: State a limitation or common misunderstanding.\n"
                f"55–60s: {cta}"
            ),
            asset_brief="Use a reviewed script, original visuals, and no music or clips without permission.",
            accessibility_note="Add accurate captions and ensure spoken explanations cover on-screen mathematics or diagrams.",
        )
    return PlatformDraft(
        platform=platform,
        format="community broadcast draft",
        headline=f"Reviewed update: {request.topic}",
        draft_copy=(
            f"Hello. This is a reviewable educational update for {request.audience}.\n\n"
            f"Verified context: {fact}\n\n{cta}\n\n"
            "Send only after the author confirms recipient consent and the final text."
        ),
        asset_brief="Use a concise text-first message; attach materials only with permission and attribution.",
        accessibility_note="Keep the message readable on mobile devices and provide text alternatives for attachments.",
    )


def build_social_media_package(
    request: SocialMediaRequest, *, provider: LLMProvider | None = None
) -> SocialMediaPackage:
    """Build reviewable drafts with bounded internal asset-planning synthesis only."""
    request.validate()
    drafts = tuple(_draft_for_platform(request, platform) for platform in request.platforms)
    drafts = _synthesise_asset_briefs(request, drafts, provider or get_provider())
    calendar = tuple(
        CalendarEntry(
            week=(index % request.calendar_weeks) + 1,
            platform=platform,
            purpose=f"Review and adapt the {platform.value} draft for {request.objective.value.replace('_', ' ')}.",
        )
        for index, platform in enumerate(request.platforms)
    )
    checks = (
        "Verify every factual, academic, credential, performance, and outcome claim against approved sources.",
        "Confirm copyright permission, licence, or fair-use basis for every external image, clip, audio item, dataset, and quotation.",
        "Check meaningful alt text, captions, readable contrast, and text alternatives for all visual material.",
        "Confirm recipient consent before any WhatsApp community or broadcast communication.",
        "Apply the 48-hour same-topic cooldown on each platform before proposing a publication time.",
        "Obtain explicit human approval before scheduling, uploading, sharing, or publishing any draft.",
        "Do not use mass direct messages, scraping, engagement bait, or unapproved account automation.",
    )
    return SocialMediaPackage(
        request=request,
        platform_drafts=drafts,
        proposed_calendar=calendar,
        publication_checks=checks,
    )



_SOCIAL_MEDIA_SYSTEM_PROMPT = """You are the Axioms internal social-media asset-planning assistant.
Improve only the supplied internal asset brief for each platform draft. Do not
write or revise post copy, captions, headlines, calls to action, publication
times, campaign plans, messages, or final public material. Preserve each
platform, platform-native format, verified facts, approved source scope,
audience, objective, brand voice, accessibility note, 48-hour cooldown, and
human approval boundary. Do not invent claims, credentials, metrics, citations,
endorsements, permissions, sources, or platform outcomes. Never schedule,
upload, publish, share, message, or authorize an external action.

Return exactly one concise internal asset-planning paragraph per supplied
platform, in its original order, separated by a line containing only `---`.
Do not use headings, lists, citations, links, or questions.
"""


def _synthesise_asset_briefs(
    request: SocialMediaRequest,
    drafts: tuple[PlatformDraft, ...],
    provider: LLMProvider,
) -> tuple[PlatformDraft, ...]:
    """Apply only well-formed internal asset guidance; all public-facing draft fields stay fixed."""
    facts = "\n".join(f"- {fact}" for fact in request.verified_facts)
    draft_specification = "\n".join(
        (
            f"{index}. Platform: {draft.platform.value}; format: {draft.format}; "
            f"current asset brief: {draft.asset_brief}; accessibility note: {draft.accessibility_note}"
        )
        for index, draft in enumerate(drafts, start=1)
    )
    user_prompt = (
        f"Topic: {request.topic}\n"
        f"Audience: {request.audience}\n"
        f"Objective: {request.objective.value}\n"
        f"Brand voice: {request.brand_voice}\n"
        f"Approved source scope: {request.approved_source_scope}\n"
        f"Verified facts:\n{facts}\n\n"
        f"Platform drafts:\n{draft_specification}\n"
    )
    try:
        result = provider.complete(
            [LLMMessage("system", _SOCIAL_MEDIA_SYSTEM_PROMPT), LLMMessage("user", user_prompt)],
            max_tokens=1_000,
            temperature=0.2,
        )
    except LLMDisabledError:
        return drafts
    except Exception:  # noqa: BLE001 - failed synthesis must retain the review-first package
        return drafts
    generated = _parse_asset_briefs(result.text, len(drafts))
    if generated is None:
        return drafts
    return tuple(
        PlatformDraft(
            platform=draft.platform,
            format=draft.format,
            headline=draft.headline,
            draft_copy=draft.draft_copy,
            asset_brief=asset_brief,
            accessibility_note=draft.accessibility_note,
        )
        for draft, asset_brief in zip(drafts, generated, strict=True)
    )


def _parse_asset_briefs(text: str, expected_drafts: int) -> tuple[str, ...] | None:
    """Reject malformed or externally actionable output before it alters internal asset guidance."""
    blocks = [block.strip() for block in text.split("---") if block.strip()]
    if len(blocks) != expected_drafts:
        return None
    if any(
        "\n" in block
        or "?" in block
        or "http://" in block.casefold()
        or "https://" in block.casefold()
        or block.startswith(("-", "*", "#"))
        or len(block) > 1_500
        for block in blocks
    ):
        return None
    return tuple(blocks)


def export_docx(package: SocialMediaPackage, destination: Path) -> Path:
    """Export a review package only; no social-platform integration is performed."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"Social Media Package: {package.request.topic}", level=0)
    document.add_paragraph(f"Audience: {package.request.audience}")
    document.add_paragraph(f"Objective: {package.request.objective.value.replace('_', ' ').title()}")
    document.add_paragraph("Status: Draft only — human approval is required before any external action.")
    for draft in package.platform_drafts:
        document.add_heading(f"{draft.platform.value.title()}: {draft.format}", level=1)
        document.add_heading(draft.headline, level=2)
        document.add_paragraph(draft.draft_copy)
        document.add_paragraph(f"Asset brief: {draft.asset_brief}")
        document.add_paragraph(f"Accessibility: {draft.accessibility_note}")
    document.add_heading("Proposed calendar", level=1)
    table = document.add_table(rows=1, cols=4)
    for cell, heading in zip(table.rows[0].cells, ("Week", "Platform", "Purpose", "Status"), strict=True):
        cell.text = heading
    for entry in package.proposed_calendar:
        cells = table.add_row().cells
        cells[0].text = str(entry.week)
        cells[1].text = entry.platform.value
        cells[2].text = entry.purpose
        cells[3].text = entry.status
    document.add_heading("Publication checks", level=1)
    for item in package.publication_checks:
        document.add_paragraph(item, style="List Bullet")
    document.save(destination)
    return destination
