"""Review-first educational content planner for videos, course modules, and workshops."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document

from axioms.llm import LLMDisabledError, LLMMessage, LLMProvider, get_provider


class ContentFormat(StrEnum):
    YOUTUBE_VIDEO = "youtube_video"
    COURSE_MODULE = "course_module"
    WORKSHOP = "workshop"


class LanguageMode(StrEnum):
    ENGLISH = "english"
    URDU = "urdu"
    BILINGUAL = "bilingual"


@dataclass(frozen=True, slots=True)
class ContentRequest:
    topic: str
    format: ContentFormat
    audience: str
    duration_minutes: int
    approved_source_scope: str
    learning_outcomes: tuple[str, ...]
    language_mode: LanguageMode = LanguageMode.ENGLISH
    application_context: str | None = None
    keywords: tuple[str, ...] = ()

    def validate(self) -> None:
        if not self.topic.strip() or not self.audience.strip() or not self.approved_source_scope.strip():
            raise ValueError("Topic, audience, and approved source scope are required.")
        if not 3 <= self.duration_minutes <= 240:
            raise ValueError("Content duration must be between 3 and 240 minutes.")
        if not 1 <= len(self.learning_outcomes) <= 6:
            raise ValueError("Provide between one and six learning outcomes.")
        if any(not outcome.strip() for outcome in self.learning_outcomes):
            raise ValueError("Learning outcomes cannot be blank.")
        if len(self.keywords) > 8 or any(not keyword.strip() for keyword in self.keywords):
            raise ValueError("Provide at most eight non-blank keywords.")


@dataclass(frozen=True, slots=True)
class ContentSegment:
    title: str
    minutes: int
    purpose: str
    speaker_notes: str


@dataclass(frozen=True, slots=True)
class ContentPackage:
    request: ContentRequest
    title_options: tuple[str, ...]
    description_framework: str
    segments: tuple[ContentSegment, ...]
    thumbnail_brief: str
    accessibility_checks: tuple[str, ...]
    accuracy_checks: tuple[str, ...]
    publication_approval_required: bool = True

    @property
    def allocated_minutes(self) -> int:
        return sum(segment.minutes for segment in self.segments)

    def to_dict(self) -> dict:
        return asdict(self)

    def to_markdown(self) -> str:
        outcomes = "\n".join(f"- {item}" for item in self.request.learning_outcomes)
        segments = "\n".join(
            f"| {item.minutes} | {item.title} | {item.purpose} |" for item in self.segments
        )
        accessibility = "\n".join(f"- [ ] {item}" for item in self.accessibility_checks)
        accuracy = "\n".join(f"- [ ] {item}" for item in self.accuracy_checks)
        titles = "\n".join(f"- {item}" for item in self.title_options)
        return f"""# Content package: {self.request.topic}

## Title options
{titles}

## Audience and learning outcomes
- Audience: {self.request.audience}
- Language mode: {self.request.language_mode.value}
{outcomes}

## Description framework
{self.description_framework}

## Script or module plan
| Minutes | Segment | Purpose |
| ---: | --- | --- |
{segments}

## Thumbnail brief
{self.thumbnail_brief}

## Accessibility checks
{accessibility}

## Accuracy and publication checks
{accuracy}
"""


def _allocate_minutes(total: int) -> tuple[int, ...]:
    weights = (0.10, 0.20, 0.30, 0.25, 0.15)
    raw = [total * weight for weight in weights]
    allocation = [max(1, int(value)) for value in raw]
    remainder = total - sum(allocation)
    fractions = sorted(
        range(len(raw)), key=lambda index: (raw[index] - int(raw[index]), -index), reverse=True
    )
    for index in fractions[:remainder]:
        allocation[index] += 1
    return tuple(allocation)


def build_content_package(
    request: ContentRequest, *, provider: LLMProvider | None = None
) -> ContentPackage:
    """Create a review-first content blueprint with bounded internal planning synthesis."""
    request.validate()
    minutes = _allocate_minutes(request.duration_minutes)
    application = request.application_context or "an instructor-approved real-world application"
    language_note = (
        "Mark each segment for English and Urdu review; translate technical terms only after subject-matter review."
        if request.language_mode is LanguageMode.BILINGUAL
        else f"Use {request.language_mode.value} with consistent, accessible terminology."
    )
    title_options = (
        f"{request.topic}: Intuition, Method, and Application",
        f"Understanding {request.topic} Through a Worked Example",
        f"{request.topic} for {request.audience}",
    )
    description = (
        f"Introduce the learning purpose, the approved scope, and who this {request.format.value.replace('_', ' ')} is for. "
        f"State only author-verified context: {request.approved_source_scope}. "
        f"Use keywords as optional metadata prompts, not as claims of search ranking: {', '.join(request.keywords) or 'none supplied'}."
    )
    segments = (
        ContentSegment("Opening hook", minutes[0], "Motivate the topic with an accurate, familiar question.", "Use an instructor-verified example; avoid sensational or clickbait wording."),
        ContentSegment("Intuition and context", minutes[1], "Connect prior knowledge to the topic.", f"Introduce {application} only after confirming its factual accuracy."),
        ContentSegment("Formal core", minutes[2], "Present definitions, assumptions, notation, or method steps precisely.", "Leave source placeholders for each non-trivial factual or mathematical statement."),
        ContentSegment("Worked application", minutes[3], "Model a complete, reviewable example or demonstration.", "Show intermediate reasoning and state any limitations or edge cases."),
        ContentSegment("Recap and next step", minutes[4], "Consolidate learning and prompt active follow-up.", f"{language_note} End with a practice prompt rather than an unverified promise."),
    )
    segments = _synthesise_speaker_notes(request, segments, provider or get_provider())
    thumbnail = (
        f"Use a clear visual of the central concept in {request.topic}, a short accurate headline, strong contrast, "
        "and no numerical performance, credential, or outcome claim unless verified by the author. Include alt-text intent."
    )
    accessibility = (
        "Provide accurate captions or a reviewed transcript before publication.",
        "Write meaningful alt text for thumbnails, diagrams, and visual demonstrations.",
        "Use readable text, high contrast, and avoid unexplained jargon.",
        "Check that equations, diagrams, and code demonstrations are also explained in speech or text.",
    )
    accuracy = (
        "Verify every mathematical, scientific, historical, and application claim against approved sources.",
        "Confirm copyright permission or fair-use basis for every third-party visual, diagram, excerpt, or dataset.",
        "Do not claim SEO rank, learner outcomes, certification, or platform performance without evidence.",
        "Obtain human approval before publishing, scheduling, uploading, or sharing any public-facing asset.",
    )
    return ContentPackage(
        request=request,
        title_options=title_options,
        description_framework=description,
        segments=segments,
        thumbnail_brief=thumbnail,
        accessibility_checks=accessibility,
        accuracy_checks=accuracy,
    )



_CONTENT_SYSTEM_PROMPT = """You are the Axioms internal content-planning assistant.
Improve only the speaker notes for the supplied deterministic content-plan
segments. These are internal editorial notes, not a final script, publication
copy, title, description, social post, factual lesson material, or publishing
plan. Preserve every segment title, duration, purpose, approved source scope,
learning outcome, language-review requirement, accessibility check, and human
publication-approval boundary. Do not invent claims, examples, citations,
sources, permissions, learner outcomes, or platform performance.

Return exactly one concise planning paragraph per supplied segment, in its
original order, separated by a line containing only `---`. Do not use headings,
lists, citations, links, or questions.
"""


def _synthesise_speaker_notes(
    request: ContentRequest,
    segments: tuple[ContentSegment, ...],
    provider: LLMProvider,
) -> tuple[ContentSegment, ...]:
    """Accept only bounded internal notes; provider failures retain the deterministic plan."""
    outcomes = "\n".join(f"- {outcome}" for outcome in request.learning_outcomes)
    segment_specification = "\n".join(
        (
            f"{index}. {segment.title}; {segment.minutes} minutes; "
            f"purpose: {segment.purpose}; current note: {segment.speaker_notes}"
        )
        for index, segment in enumerate(segments, start=1)
    )
    user_prompt = (
        f"Topic: {request.topic}\n"
        f"Format: {request.format.value}\n"
        f"Audience: {request.audience}\n"
        f"Language mode: {request.language_mode.value}\n"
        f"Approved source scope: {request.approved_source_scope}\n"
        f"Learning outcomes:\n{outcomes}\n\n"
        f"Content-plan segments:\n{segment_specification}\n"
    )
    try:
        result = provider.complete(
            [LLMMessage("system", _CONTENT_SYSTEM_PROMPT), LLMMessage("user", user_prompt)],
            max_tokens=1_200,
            temperature=0.2,
        )
    except LLMDisabledError:
        return segments
    except Exception:  # noqa: BLE001 - failed synthesis must not weaken review-first planning
        return segments
    generated = _parse_speaker_notes(result.text, len(segments))
    if generated is None:
        return segments
    return tuple(
        ContentSegment(
            title=segment.title,
            minutes=segment.minutes,
            purpose=segment.purpose,
            speaker_notes=speaker_notes,
        )
        for segment, speaker_notes in zip(segments, generated, strict=True)
    )


def _parse_speaker_notes(text: str, expected_segments: int) -> tuple[str, ...] | None:
    """Reject malformed, list-like, link-like, or question-like output before it alters the plan."""
    blocks = [block.strip() for block in text.split("---") if block.strip()]
    if len(blocks) != expected_segments:
        return None
    if any(
        "\n" in block
        or "?" in block
        or "http://" in block.casefold()
        or "https://" in block.casefold()
        or block.startswith(("-", "*", "#"))
        or len(block) > 1_500
        for block in blocks
    ):
        return None
    return tuple(blocks)


def export_docx(package: ContentPackage, destination: Path) -> Path:
    """Export a reviewable content plan; this function does not publish any content."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"Content Package: {package.request.topic}", level=0)
    document.add_paragraph(f"Format: {package.request.format.value.replace('_', ' ').title()}")
    document.add_paragraph(f"Audience: {package.request.audience} | Language: {package.request.language_mode.value}")
    document.add_heading("Title options", level=1)
    for title in package.title_options:
        document.add_paragraph(title, style="List Bullet")
    document.add_heading("Description framework", level=1)
    document.add_paragraph(package.description_framework)
    document.add_heading("Script or module plan", level=1)
    table = document.add_table(rows=1, cols=3)
    for cell, heading in zip(table.rows[0].cells, ("Minutes", "Segment", "Purpose"), strict=True):
        cell.text = heading
    for segment in package.segments:
        cells = table.add_row().cells
        cells[0].text = str(segment.minutes)
        cells[1].text = segment.title
        cells[2].text = segment.purpose
    document.add_heading("Thumbnail brief", level=1)
    document.add_paragraph(package.thumbnail_brief)
    document.add_heading("Accessibility checks", level=1)
    for item in package.accessibility_checks:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Accuracy and publication checks", level=1)
    for item in package.accuracy_checks:
        document.add_paragraph(item, style="List Bullet")
    document.save(destination)
    return destination

