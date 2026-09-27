"""Evidence-first research agent core with explicit source and claim verification states."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document


class PublicationKind(StrEnum):
    JOURNAL_ARTICLE = "journal_article"
    CONFERENCE_PAPER = "conference_paper"
    PREPRINT = "preprint"
    BOOK = "book"
    OTHER = "other"


class VerificationStatus(StrEnum):
    UNVERIFIED = "unverified"
    METADATA_VERIFIED = "metadata_verified"
    CLAIM_VERIFIED = "claim_verified"
    REJECTED = "rejected"


DOI_PATTERN = re.compile(r"^10\.\d{4,9}/[-._;()/:a-z0-9]+$", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class EvidenceSource:
    source_id: str
    title: str
    authors: tuple[str, ...]
    year: int
    publication_kind: PublicationKind
    doi: str | None
    url: str | None
    peer_reviewed: bool
    supported_claim: str | None
    verification_status: VerificationStatus
    verification_evidence: str | None = None

    def normalized_doi(self) -> str | None:
        if self.doi is None:
            return None
        return self.doi.removeprefix("https://doi.org/").removeprefix("http://doi.org/").strip()

    def validate(self) -> None:
        if not self.source_id.strip() or not self.title.strip() or not self.authors:
            raise ValueError("Each source needs an ID, title, and at least one author.")
        if not 1600 <= self.year <= 2100:
            raise ValueError("Source year must be between 1600 and 2100.")
        if not self.doi and not self.url:
            raise ValueError("Each source needs a DOI or a source URL.")
        doi = self.normalized_doi()
        if doi and not DOI_PATTERN.fullmatch(doi):
            raise ValueError(f"Source {self.source_id} has an invalid DOI format.")
        if self.peer_reviewed and self.publication_kind is PublicationKind.PREPRINT:
            raise ValueError("A preprint cannot be labelled peer reviewed without a separate published record.")
        if self.verification_status is not VerificationStatus.UNVERIFIED and not self.verification_evidence:
            raise ValueError("A non-unverified source needs a verification-evidence note.")
        if self.verification_status is VerificationStatus.CLAIM_VERIFIED:
            if not doi:
                raise ValueError("Claim-verified sources require a DOI in this MVP.")
            if not self.supported_claim or not self.supported_claim.strip():
                raise ValueError("Claim-verified sources require the supported claim to be recorded.")

    def citation_label(self) -> str:
        author = self.authors[0] if len(self.authors) == 1 else f"{self.authors[0]} et al."
        return f"{author} ({self.year}) — {self.title}"


@dataclass(frozen=True, slots=True)
class ResearchRequest:
    research_question: str
    scope: str
    sources: tuple[EvidenceSource, ...]
    analysis_dimensions: tuple[str, ...] = ("methods", "datasets", "metrics", "limitations")
    target_venue: str | None = None

    def validate(self) -> None:
        if not self.research_question.strip() or not self.scope.strip():
            raise ValueError("Research question and scope are required.")
        if not self.sources:
            raise ValueError("Provide at least one source record before synthesis.")
        source_ids = [source.source_id for source in self.sources]
        if len(source_ids) != len(set(source_ids)):
            raise ValueError("Source IDs must be unique.")
        if not self.analysis_dimensions or any(not item.strip() for item in self.analysis_dimensions):
            raise ValueError("Provide at least one non-blank analysis dimension.")
        for source in self.sources:
            source.validate()


@dataclass(frozen=True, slots=True)
class SourceAudit:
    source_id: str
    citation: str
    verification_status: VerificationStatus
    synthesis_eligible: bool
    finding: str


@dataclass(frozen=True, slots=True)
class ResearchBrief:
    request: ResearchRequest
    source_audit: tuple[SourceAudit, ...]
    verified_claims: tuple[str, ...]
    blocked_claims: tuple[str, ...]
    verification_queue: tuple[str, ...]
    synthesis_boundary: str
    human_review_required: bool = True

    def to_dict(self) -> dict:
        return asdict(self)

    def to_markdown(self) -> str:
        audit_rows = "\n".join(
            f"| {item.source_id} | {item.verification_status.value} | {'Yes' if item.synthesis_eligible else 'No'} | {item.finding} |"
            for item in self.source_audit
        )
        verified = "\n".join(f"- {claim}" for claim in self.verified_claims) or "- None yet."
        blocked = "\n".join(f"- {claim}" for claim in self.blocked_claims) or "- None."
        queue = "\n".join(f"- {item}" for item in self.verification_queue) or "- No pending checks."
        return f"""# Research brief

## Research question
{self.request.research_question}

## Scope
{self.request.scope}

## Analysis dimensions
{', '.join(self.request.analysis_dimensions)}

## Source audit
| Source | Verification state | Eligible for synthesis | Finding |
| --- | --- | --- | --- |
{audit_rows}

## Claims eligible for constrained synthesis
{verified}

## Claims blocked from synthesis
{blocked}

## Verification queue
{queue}

## Synthesis boundary
{self.synthesis_boundary}
"""

    def verified_bibtex(self) -> str:
        records = []
        for source in self.request.sources:
            if source.verification_status not in {
                VerificationStatus.METADATA_VERIFIED,
                VerificationStatus.CLAIM_VERIFIED,
            }:
                continue
            key = re.sub(r"[^a-z0-9]", "", source.authors[0].casefold()) + str(source.year)
            record_type = "article" if source.publication_kind is PublicationKind.JOURNAL_ARTICLE else "misc"
            fields = [
                f"  author = {{{' and '.join(source.authors)}}}",
                f"  title = {{{source.title}}}",
                f"  year = {{{source.year}}}",
            ]
            if source.normalized_doi():
                fields.append(f"  doi = {{{source.normalized_doi()}}}")
            if source.url:
                fields.append(f"  url = {{{source.url}}}")
            records.append(f"@{record_type}{{{key},\n" + ",\n".join(fields) + "\n}")
        return "\n\n".join(records) + ("\n" if records else "")


def build_research_brief(request: ResearchRequest) -> ResearchBrief:
    """Create an auditable evidence ledger; never infer verification from source identity alone."""
    request.validate()
    source_audit: list[SourceAudit] = []
    verified_claims: list[str] = []
    blocked_claims: list[str] = []
    verification_queue: list[str] = []
    for source in request.sources:
        eligible = source.verification_status is VerificationStatus.CLAIM_VERIFIED
        if eligible:
            verified_claims.append(f"[{source.source_id}] {source.supported_claim}")
            finding = "Claim and metadata are recorded as verified; eligible for constrained synthesis."
        elif source.verification_status is VerificationStatus.REJECTED:
            blocked_claims.append(f"[{source.source_id}] Rejected source: do not cite or synthesize its claim.")
            finding = "Rejected; excluded from citations and synthesis."
        else:
            claim = source.supported_claim or "No source-linked claim recorded."
            blocked_claims.append(f"[{source.source_id}] {claim}")
            required = "Verify metadata and DOI" if source.verification_status is VerificationStatus.UNVERIFIED else "Verify the specific claim against the source"
            verification_queue.append(f"[{source.source_id}] {required}: {source.citation_label()}")
            finding = "Not eligible until claim-level verification is recorded."
        if source.publication_kind is PublicationKind.PREPRINT:
            verification_queue.append(f"[{source.source_id}] Record preprint status; do not describe it as peer reviewed.")
        source_audit.append(
            SourceAudit(
                source_id=source.source_id,
                citation=source.citation_label(),
                verification_status=source.verification_status,
                synthesis_eligible=eligible,
                finding=finding,
            )
        )
    boundary = (
        "This brief may synthesize only the claim-verified records shown above. "
        "It must abstain from factual conclusions, novelty claims, comparative performance claims, "
        "or recommendations that depend on unverified or rejected records."
    )
    return ResearchBrief(
        request=request,
        source_audit=tuple(source_audit),
        verified_claims=tuple(verified_claims),
        blocked_claims=tuple(blocked_claims),
        verification_queue=tuple(verification_queue),
        synthesis_boundary=boundary,
    )


def export_docx(brief: ResearchBrief, destination: Path) -> Path:
    """Export the evidence ledger and its limits for human review."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Research Brief", level=0)
    document.add_heading("Research question", level=1)
    document.add_paragraph(brief.request.research_question)
    document.add_heading("Scope", level=1)
    document.add_paragraph(brief.request.scope)
    document.add_heading("Source audit", level=1)
    table = document.add_table(rows=1, cols=4)
    for cell, heading in zip(
        table.rows[0].cells,
        ("Source", "Status", "Synthesis eligible", "Finding"),
        strict=True,
    ):
        cell.text = heading
    for item in brief.source_audit:
        cells = table.add_row().cells
        cells[0].text = item.source_id
        cells[1].text = item.verification_status.value
        cells[2].text = "Yes" if item.synthesis_eligible else "No"
        cells[3].text = item.finding
    document.add_heading("Claims eligible for constrained synthesis", level=1)
    for claim in brief.verified_claims or ("None yet.",):
        document.add_paragraph(claim, style="List Bullet")
    document.add_heading("Verification queue", level=1)
    for item in brief.verification_queue or ("No pending checks.",):
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Synthesis boundary", level=1)
    document.add_paragraph(brief.synthesis_boundary)
    document.save(destination)
    return destination

