"""Evidence-bound STEM AI portfolio planning without external repository operations."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document

from axioms.llm import LLMDisabledError, LLMMessage, LLMProvider, get_provider


class PortfolioAudience(StrEnum):
    ACADEMIC = "academic"
    INDUSTRY = "industry"
    BOTH = "academic_and_industry"


class RepositoryVisibility(StrEnum):
    PRIVATE = "private"
    PUBLIC = "public"


class DataAccessLevel(StrEnum):
    PUBLIC = "public"
    DE_IDENTIFIED = "de_identified"
    PROPRIETARY = "proprietary"
    RESTRICTED = "restricted"
    UNKNOWN = "unknown"


@dataclass(frozen=True, slots=True)
class PortfolioEvidence:
    evidence_id: str
    claim: str
    source_reference: str
    verified: bool

    def validate(self) -> None:
        if not self.evidence_id.strip() or not self.claim.strip() or not self.source_reference.strip():
            raise ValueError("Each portfolio evidence record needs an ID, claim, and source reference.")


@dataclass(frozen=True, slots=True)
class DatasetAsset:
    name: str
    access_level: DataAccessLevel
    licence_or_permission: str | None = None
    attribution: str | None = None

    def validate(self) -> None:
        if not self.name.strip():
            raise ValueError("Dataset asset names cannot be blank.")
        if (
            self.access_level in {DataAccessLevel.PUBLIC, DataAccessLevel.DE_IDENTIFIED}
            and (not self.licence_or_permission or not self.attribution)
        ):
            raise ValueError(
                "Public or de-identified dataset assets need licence/permission and attribution details."
            )


@dataclass(frozen=True, slots=True)
class PortfolioRequest:
    project_title: str
    research_summary: str
    target_audience: PortfolioAudience
    repository_visibility: RepositoryVisibility
    verified_evidence: tuple[PortfolioEvidence, ...]
    dataset_assets: tuple[DatasetAsset, ...] = ()
    include_demo_plan: bool = True
    include_notebook_plan: bool = True

    def validate(self) -> None:
        if not self.project_title.strip() or not self.research_summary.strip():
            raise ValueError("Project title and research summary are required.")
        if not 1 <= len(self.verified_evidence) <= 12:
            raise ValueError("Provide between one and twelve evidence records.")
        evidence_ids = [item.evidence_id for item in self.verified_evidence]
        if len(evidence_ids) != len(set(evidence_ids)):
            raise ValueError("Portfolio evidence IDs must be unique.")
        for item in self.verified_evidence:
            item.validate()
            if not item.verified:
                raise ValueError("Only verified evidence may be included in a portfolio blueprint.")
        for asset in self.dataset_assets:
            asset.validate()
        if self.repository_visibility is RepositoryVisibility.PUBLIC:
            blocked_assets = [
                asset.name
                for asset in self.dataset_assets
                if asset.access_level
                in {DataAccessLevel.PROPRIETARY, DataAccessLevel.RESTRICTED, DataAccessLevel.UNKNOWN}
            ]
            if blocked_assets:
                raise ValueError(
                    "A public repository cannot include proprietary, restricted, or unknown-access assets: "
                    + ", ".join(blocked_assets)
                )


@dataclass(frozen=True, slots=True)
class RepositoryItem:
    path: str
    purpose: str


@dataclass(frozen=True, slots=True)
class ReadmeSection:
    heading: str
    guidance: str


@dataclass(frozen=True, slots=True)
class PortfolioPackage:
    request: PortfolioRequest
    repository_structure: tuple[RepositoryItem, ...]
    readme_sections: tuple[ReadmeSection, ...]
    reproducibility_checklist: tuple[str, ...]
    impact_summary: tuple[str, ...]
    publication_checks: tuple[str, ...]
    approval_required: bool = True
    github_action_blocked: bool = True

    def to_dict(self) -> dict:
        return asdict(self)

    def to_markdown(self) -> str:
        structure = "\n".join(f"- `{item.path}` — {item.purpose}" for item in self.repository_structure)
        readme = "\n".join(
            f"## {item.heading}\n{item.guidance}" for item in self.readme_sections
        )
        reproducibility = "\n".join(f"- [ ] {item}" for item in self.reproducibility_checklist)
        impact = "\n".join(f"- {item}" for item in self.impact_summary)
        checks = "\n".join(f"- [ ] {item}" for item in self.publication_checks)
        return f"""# Portfolio package: {self.request.project_title}

## Research summary
{self.request.research_summary}

## Repository structure
{structure}

## README plan
{readme}

## Reproducibility checklist
{reproducibility}

## Evidence-bound impact summary
{impact}

## Publication checks
{checks}
"""


def build_portfolio_package(
    request: PortfolioRequest, *, provider: LLMProvider | None = None
) -> PortfolioPackage:
    """Create an evidence-bound repository plan with bounded internal structure synthesis."""
    request.validate()
    structure = [
        RepositoryItem("README.md", "Scope, verified claims, installation, reproduction, and limitations."),
        RepositoryItem("LICENSE", "Chosen and confirmed licence for the repository contents."),
        RepositoryItem("pyproject.toml", "Pinned Python project metadata and runtime dependencies."),
        RepositoryItem("src/", "Implementation modules with clear public interfaces."),
        RepositoryItem("tests/", "Automated unit, boundary, negative, and regression tests."),
        RepositoryItem("docs/methodology.md", "Methods, assumptions, and evidence-to-claim traceability."),
        RepositoryItem("docs/data_provenance.md", "Dataset source, access, licence, attribution, and exclusions."),
        RepositoryItem("scripts/reproduce.py", "Deterministic reproduction entry point and environment checks."),
        RepositoryItem(".github/workflows/ci.yml", "Test and lint workflow; no secrets or deployment action."),
    ]
    if request.include_notebook_plan:
        structure.append(
            RepositoryItem("notebooks/01_reproducible_demo.ipynb", "A documented, rerunnable demonstration plan.")
        )
    if request.include_demo_plan:
        structure.append(
            RepositoryItem("demo/", "A local, dry-run demonstration; no public hosting assumption."))
    structure = list(_synthesise_repository_purposes(request, tuple(structure), provider or get_provider()))
    evidence_lines = tuple(
        f"[{item.evidence_id}] {item.claim} (source: {item.source_reference})"
        for item in request.verified_evidence
    )
    readme_sections = (
        ReadmeSection("Project purpose", request.research_summary),
        ReadmeSection(
            "Evidence-bound claims",
            "List only the following author-verified statements:\n" + "\n".join(evidence_lines),
        ),
        ReadmeSection(
            "Reproduction", "State the pinned environment, data-access steps, command sequence, and expected outputs.",
        ),
        ReadmeSection(
            "Limitations and responsible use",
            "State assumptions, data-access boundaries, failure modes, and claims that the project does not establish.",
        ),
    )
    reproducibility = (
        "Pin runtime and development dependencies; record the interpreter version.",
        "Provide a deterministic seed and a documented command to reproduce every reported artifact.",
        "Run tests, linting, and an end-to-end dry run before publication.",
        "Record dataset version, licence, attribution, preprocessing, and exclusion criteria.",
        "Keep raw restricted, proprietary, personal, or unknown-access data outside the public repository.",
        "Document expected outputs and known non-determinism without implying a performance guarantee.",
    )
    impact = (
        "Portfolio claims are constrained to the verified evidence records supplied by the author.",
        "No GitHub-star, citation, H-index, acceptance, novelty, or performance outcome is predicted or inferred.",
        "The package is a reproducibility and communication plan, not evidence that code, data, or results are publication-ready.",
    )
    checks = (
        "Confirm that every public claim maps to a verified evidence ID and source reference.",
        "Confirm all code has passed its intended tests and that the reported environment can reproduce the documented outputs.",
        "Confirm dataset licences, permissions, attribution, and public-release eligibility.",
        "Remove API keys, credentials, student information, unpublished results, and proprietary or restricted data.",
        "Obtain explicit human approval before creating, pushing to, or changing any GitHub repository.",
    )
    return PortfolioPackage(
        request=request,
        repository_structure=tuple(structure),
        readme_sections=readme_sections,
        reproducibility_checklist=reproducibility,
        impact_summary=impact,
        publication_checks=checks,
    )



_PORTFOLIO_SYSTEM_PROMPT = """You are the Axioms internal STEM AI portfolio-planning assistant.
Improve only the internal purpose description for each supplied repository-plan
item. Do not write a README, source code, notebook, demo, documentation,
command, repository content, public portfolio copy, deployment plan, or GitHub
action. Preserve every repository path, verified evidence record, dataset
access boundary, requested visibility, licence and attribution requirement,
testing-before-publication requirement, reproducibility requirement, and human
approval boundary. Do not invent claims, citations, performance values,
research outcomes, licences, permissions, data access, credentials, GitHub
metadata, external links, novelty, or impact. Never create, change, commit,
push, deploy, publish, or authorize a repository action.

Return exactly one concise internal planning paragraph per supplied repository
item, in its original order, separated by a line containing only `---`. Do not
use headings, lists, code, commands, citations, links, or questions.
"""


def _synthesise_repository_purposes(
    request: PortfolioRequest,
    structure: tuple[RepositoryItem, ...],
    provider: LLMProvider,
) -> tuple[RepositoryItem, ...]:
    """Use only well-formed internal descriptions; all evidence and release controls stay fixed."""
    evidence = "\n".join(
        f"- [{item.evidence_id}] {item.claim} (source: {item.source_reference})"
        for item in request.verified_evidence
    )
    assets = "\n".join(
        f"- {asset.name}: {asset.access_level.value}; permission: {asset.licence_or_permission or 'not supplied'}; "
        f"attribution: {asset.attribution or 'not supplied'}"
        for asset in request.dataset_assets
    ) or "- No dataset assets declared."
    plan_items = "\n".join(
        f"{index}. Path: {item.path}; current purpose: {item.purpose}"
        for index, item in enumerate(structure, start=1)
    )
    user_prompt = (
        f"Project title: {request.project_title}\n"
        f"Audience: {request.target_audience.value}\n"
        f"Proposed repository visibility: {request.repository_visibility.value}\n"
        f"Research summary: {request.research_summary}\n\n"
        f"Verified evidence records:\n{evidence}\n\n"
        f"Declared dataset assets:\n{assets}\n\n"
        f"Repository-plan items:\n{plan_items}\n"
    )
    try:
        result = provider.complete(
            [LLMMessage("system", _PORTFOLIO_SYSTEM_PROMPT), LLMMessage("user", user_prompt)],
            max_tokens=1_500,
            temperature=0.2,
        )
    except LLMDisabledError:
        return structure
    except Exception:  # noqa: BLE001 - failed synthesis must retain the evidence-bound package
        return structure
    generated = _parse_repository_purposes(result.text, len(structure))
    if generated is None:
        return structure
    return tuple(
        RepositoryItem(path=item.path, purpose=purpose)
        for item, purpose in zip(structure, generated, strict=True)
    )


def _parse_repository_purposes(text: str, expected_items: int) -> tuple[str, ...] | None:
    """Reject malformed or externally actionable output before it changes the internal plan."""
    blocks = [block.strip() for block in text.split("---") if block.strip()]
    if len(blocks) != expected_items:
        return None
    if any(
        "\n" in block
        or "?" in block
        or "http://" in block.casefold()
        or "https://" in block.casefold()
        or block.startswith(("-", "*", "#", "`"))
        or len(block) > 1_500
        for block in blocks
    ):
        return None
    return tuple(blocks)


def export_docx(package: PortfolioPackage, destination: Path) -> Path:
    """Export a portfolio-review document without creating or updating external repositories."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(f"Portfolio Package: {package.request.project_title}", level=0)
    document.add_paragraph(f"Audience: {package.request.target_audience.value.replace('_', ' ')}")
    document.add_paragraph(f"Proposed repository visibility: {package.request.repository_visibility.value}")
    document.add_paragraph("Status: Review required — GitHub actions are not performed by this agent.")
    document.add_heading("Repository structure", level=1)
    for item in package.repository_structure:
        document.add_paragraph(f"{item.path}: {item.purpose}", style="List Bullet")
    document.add_heading("README plan", level=1)
    for item in package.readme_sections:
        document.add_heading(item.heading, level=2)
        document.add_paragraph(item.guidance)
    document.add_heading("Reproducibility checklist", level=1)
    for item in package.reproducibility_checklist:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Evidence-bound impact summary", level=1)
    for item in package.impact_summary:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Publication checks", level=1)
    for item in package.publication_checks:
        document.add_paragraph(item, style="List Bullet")
    document.save(destination)
    return destination
